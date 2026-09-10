"""Tests for two connection-related features requested directly: a
custom/free-text Purpose for a connection (the example given: "ILO/XCC",
a server management interface not covered by the fixed Uplink/Data/
Storage/Management/vMotion/Other list), and an optional "Notes" column
in the Word report's Connections table - the user had put "ILO/XCC"
into a connection's Notes as a workaround, then found it never showed
up in the exported report at all.
"""

import pytest

pytest.importorskip("PySide6")

from unittest.mock import patch

from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from src.calculations.docx_report import build_docx_report
from src.calculations.thresholds import Thresholds
from src.models.cluster_project import ClusterProject
from src.models.network_connection import NetworkConnection
from src.models.network_switch import NetworkSwitch
from src.models.server import Server


@pytest.fixture(scope="module", autouse=True)
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


# ----------------------------------------------------------------------
# Editable Purpose combo
# ----------------------------------------------------------------------

def test_purpose_combo_is_editable():
    from src.gui.dialogs.connection_dialog import ConnectionDialog

    project = ClusterProject()
    dialog = ConnectionDialog(project)

    assert dialog.purpose_combo.isEditable() is True


def test_purpose_combo_still_lists_the_standard_options():
    from src.gui.dialogs.connection_dialog import ConnectionDialog
    from src.models.network_connection import PURPOSE_OPTIONS

    project = ClusterProject()
    dialog = ConnectionDialog(project)

    items = [dialog.purpose_combo.itemText(i) for i in range(dialog.purpose_combo.count())]
    for option in PURPOSE_OPTIONS:
        assert option in items


def test_custom_purpose_saves_correctly():
    from src.gui.dialogs.connection_dialog import ConnectionDialog

    project = ClusterProject()
    srv = Server.create_default()
    sw = NetworkSwitch.create_default()
    project.servers.append(srv)
    project.switches.append(sw)
    dialog = ConnectionDialog(project)
    dialog.combo_a.setCurrentIndex(dialog.combo_a.findData(srv.uid))
    dialog.combo_b.setCurrentIndex(dialog.combo_b.findData(sw.uid))
    dialog.purpose_combo.setCurrentText("ILO/XCC")

    connection = dialog.get_connection()

    assert connection.purpose == "ILO/XCC"


def test_editing_an_existing_custom_purpose_preloads_it():
    from src.gui.dialogs.connection_dialog import ConnectionDialog

    project = ClusterProject()
    conn = NetworkConnection.create_default()
    conn.purpose = "ILO/XCC"

    dialog = ConnectionDialog(project, conn)

    assert dialog.purpose_combo.currentText() == "ILO/XCC"


def test_custom_purpose_clsz_round_trip(tmp_path):
    from src.persistence import project_repository

    project = ClusterProject(name="Custom purpose round trip")
    conn = NetworkConnection.create_default()
    conn.purpose = "ILO/XCC"
    project.connections.append(conn)
    path = tmp_path / "p.clsz"

    project_repository.save_project(project, path, Thresholds())
    loaded = project_repository.load_project(path)

    assert loaded.project.connections[0].purpose == "ILO/XCC"


# ----------------------------------------------------------------------
# Optional connection Notes export
# ----------------------------------------------------------------------

def _project_with_a_noted_connection():
    project = ClusterProject()
    srv = Server.create_default()
    sw = NetworkSwitch.create_default()
    project.servers.append(srv)
    project.switches.append(sw)
    conn = NetworkConnection.create_default()
    conn.server_uid = srv.uid
    conn.switch_uid = sw.uid
    conn.purpose = "ILO/XCC"
    conn.notes = "ILO/XCC management interface"
    project.connections.append(conn)
    return project


def _connections_table(doc):
    return next(t for t in doc.tables if "Purpose" in [c.text for c in t.rows[0].cells])


def test_notes_column_absent_by_default():
    project = _project_with_a_noted_connection()

    doc = build_docx_report(project, Thresholds(), app_version="test")

    headers = [c.text for c in _connections_table(doc).rows[0].cells]
    assert "Notes" not in headers


def test_notes_column_present_when_opted_in():
    project = _project_with_a_noted_connection()

    doc = build_docx_report(project, Thresholds(), app_version="test", include_connection_notes=True)

    headers = [c.text for c in _connections_table(doc).rows[0].cells]
    assert "Notes" in headers


def test_notes_content_correct_when_included():
    project = _project_with_a_noted_connection()

    doc = build_docx_report(project, Thresholds(), app_version="test", include_connection_notes=True)

    table = _connections_table(doc)
    headers = [c.text for c in table.rows[0].cells]
    notes_col = headers.index("Notes")
    assert table.rows[1].cells[notes_col].text == "ILO/XCC management interface"


def test_notes_shows_dash_for_a_connection_with_no_notes():
    project = ClusterProject()
    srv = Server.create_default()
    sw = NetworkSwitch.create_default()
    project.servers.append(srv)
    project.switches.append(sw)
    conn = NetworkConnection.create_default()
    conn.server_uid = srv.uid
    conn.switch_uid = sw.uid
    project.connections.append(conn)

    doc = build_docx_report(project, Thresholds(), app_version="test", include_connection_notes=True)

    table = _connections_table(doc)
    headers = [c.text for c in table.rows[0].cells]
    notes_col = headers.index("Notes")
    assert table.rows[1].cells[notes_col].text == "-"


def test_reports_page_has_the_checkbox_unchecked_by_default():
    from src.gui.pages.reports_page import ReportsPage
    from src.services.project_service import ProjectService

    service = ProjectService()
    page = ReportsPage(service)

    assert page.include_connection_notes_check.isChecked() is False


def test_reports_page_checkbox_controls_the_exported_file(tmp_path):
    from docx import Document

    from src.gui.pages.reports_page import ReportsPage
    from src.services.project_service import ProjectService

    service = ProjectService()
    srv = Server.create_default()
    sw = NetworkSwitch.create_default()
    service.add_server(srv)
    service.add_switch(sw)
    conn = NetworkConnection.create_default()
    conn.server_uid = srv.uid
    conn.switch_uid = sw.uid
    conn.notes = "ILO/XCC management interface"
    service.project.connections.append(conn)
    page = ReportsPage(service)
    page.include_connection_notes_check.setChecked(True)
    docpath = str(tmp_path / "test.docx")

    with patch.object(QFileDialog, "getSaveFileName", return_value=(docpath, "")), \
         patch.object(QMessageBox, "information"):
        page._export_docx()

    doc = Document(docpath)
    headers = [c.text for c in _connections_table(doc).rows[0].cells]
    assert "Notes" in headers


def test_reports_page_checkbox_unchecked_excludes_notes_from_export(tmp_path):
    from docx import Document

    from src.gui.pages.reports_page import ReportsPage
    from src.services.project_service import ProjectService

    service = ProjectService()
    srv = Server.create_default()
    sw = NetworkSwitch.create_default()
    service.add_server(srv)
    service.add_switch(sw)
    conn = NetworkConnection.create_default()
    conn.server_uid = srv.uid
    conn.switch_uid = sw.uid
    conn.notes = "should not appear"
    service.project.connections.append(conn)
    page = ReportsPage(service)
    docpath = str(tmp_path / "test.docx")

    with patch.object(QFileDialog, "getSaveFileName", return_value=(docpath, "")), \
         patch.object(QMessageBox, "information"):
        page._export_docx()

    doc = Document(docpath)
    headers = [c.text for c in _connections_table(doc).rows[0].cells]
    assert "Notes" not in headers
